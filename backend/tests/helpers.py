"""Test helper functions."""

from io import BytesIO
from xml.sax.saxutils import escape

from sqlalchemy.ext.asyncio import AsyncSession

from app.auth import hash_password
from app.models import User


async def create_user(
    session: AsyncSession,
    email: str = "user@example.com",
    password: str = "password123",
    role: str = "user",
) -> User:
    """Create and persist a user, returning the ORM instance."""
    user = User(email=email, password_hash=hash_password(password), role=role)
    session.add(user)
    await session.commit()
    await session.refresh(user)
    return user


def make_docx(*paragraphs: str) -> bytes:
    """Build a minimal DOCX from plain paragraphs (used by CV parser tests).

    A paragraph starting with ``"- "`` is added with the built-in ``List Bullet``
    style — a real Word bullet whose numbering lives on the style rather than the
    run text — so tests exercise the DOCX list-item path.
    """
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        if paragraph.startswith("- "):
            document.add_paragraph(paragraph[2:], style="List Bullet")
        else:
            document.add_paragraph(paragraph)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_odt(*paragraphs: str) -> bytes:
    """Build a minimal ODT (ZIP with ``content.xml``) for CV parser tests.

    Mirrors a real ODT's structure closely enough for the stdlib extractor:
    a ``mimetype`` entry plus a ``content.xml`` holding one ``text:p`` per
    paragraph. A paragraph starting with ``"- "`` is emitted as a real
    ``text:list``/``text:list-item`` bullet (as LibreOffice does) so tests
    exercise the ODT list-item path. Paragraph text is XML-escaped.
    """
    import zipfile

    blocks: list[str] = []
    list_items: list[str] = []

    def flush_list() -> None:
        if list_items:
            items = "".join(
                f"<text:list-item><text:p>{item}</text:p></text:list-item>" for item in list_items
            )
            blocks.append(f"<text:list>{items}</text:list>")
            list_items.clear()

    for paragraph in paragraphs:
        if paragraph.startswith("- "):
            list_items.append(escape(paragraph[2:]))
            continue
        flush_list()
        blocks.append(f"<text:p>{escape(paragraph)}</text:p>")
    flush_list()

    body = "".join(blocks)
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        "<office:document-content "
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'office:version="1.3">'
        "<office:body><office:text>"
        f"{body}"
        "</office:text></office:body>"
        "</office:document-content>"
    )
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", content)
    return buffer.getvalue()


def make_pdf(*lines: str) -> bytes:
    """Build a minimal one-page PDF with one ASCII text line per row.

    Each line is its own BT/ET block with a decreasing y-coordinate so pypdf
    extracts line breaks. ASCII only: the standard Helvetica base font has no
    unicode encoding table, so non-ASCII glyphs would come out garbled.
    """
    parts = []
    for index, line in enumerate(lines):
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        parts.append(f"BT /F1 12 Tf 72 {720 - index * 14} Td ({escaped}) Tj ET")
    stream = "\n".join(parts).encode()

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % index + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_pos,
    )
    return bytes(out)
