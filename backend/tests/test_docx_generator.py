"""Tests for DOCX export — the editable sibling of the PDF (ADR 017)."""

import zipfile
from io import BytesIO

from docx import Document
from docx.shared import Pt

from app.services.cv_labels import labels_for
from app.services.cv_tailor import TailoredCV, tailor_cv
from app.services.docx_generator import generate_docx

CV = {
    "full_name": "Jane Doe",
    "email": "jane@example.com",
    "phone": "+55 11 91234-5678",
    "location": "Sao Paulo, Brazil",
    "links": ["https://github.com/janedoe"],
    "summary": "Backend engineer",
    "experience": [
        {
            "title": "Python Developer",
            "company": "Acme",
            "start_date": "2020-01",
            "end_date": None,
            "bullets": ["Built FastAPI services"],
        }
    ],
    "education": [{"institution": "MIT", "degree": "BSc", "year": "2015"}],
    "skills": ["python", "fastapi"],
    "projects": [{"name": "PudimJobs", "description": "Job tracker", "link": "https://pd.dev"}],
    "certifications": ["AWS Solutions Architect"],
    "languages": ["English (fluent)"],
}


def _document(tailored, **kwargs) -> Document:
    return Document(BytesIO(generate_docx(tailored, **kwargs)))


def _texts(document: Document) -> list[str]:
    return [paragraph.text for paragraph in document.paragraphs]


def _style_for(document: Document, text: str) -> str:
    matches = [p for p in document.paragraphs if p.text == text]
    assert matches, f"no paragraph with text {text!r}"
    return matches[0].style.name


def test_generate_docx_returns_a_valid_docx_package():
    docx = generate_docx(tailor_cv(CV, ["python"]))
    assert docx[:2] == b"PK"  # every OOXML file is a zip archive
    names = zipfile.ZipFile(BytesIO(docx)).namelist()
    assert "word/document.xml" in names
    assert "[Content_Types].xml" in names


def test_generate_docx_renders_the_sections_and_uses_word_styles():
    document = _document(tailor_cv(CV, ["python"]))

    assert _style_for(document, "Jane Doe") == "Title"
    assert _style_for(document, "Experience") == "Heading 1"
    assert _style_for(document, "Built FastAPI services") == "List Bullet"


def test_generate_docx_contact_is_one_centred_line():
    """Every field on one centred line, as the PDF header draws it (ADR 021)."""
    document = _document(tailor_cv(CV, ["python"]))
    texts = _texts(document)

    line = "jane@example.com • +55 11 91234-5678 • Sao Paulo, Brazil • https://github.com/janedoe"
    assert line in texts
    # The details/links split is gone, so neither half is a paragraph of its own.
    assert "jane@example.com • +55 11 91234-5678 • Sao Paulo, Brazil" not in texts
    assert "https://github.com/janedoe" not in texts
    assert all(paragraph.alignment is not None for paragraph in document.paragraphs[:2])
    # 8pt is the size the PDF header fits five fields at (ADR 021).
    assert document.paragraphs[1].runs[0].font.size == Pt(8)


def test_generate_docx_uses_the_pdf_page_box():
    """A4 with the PDF's side margins: Letter's 432pt is too narrow for the header."""
    section = _document(tailor_cv(CV, ["python"])).sections[0]

    # Word stores lengths in twips, so a cm round-trips to within a fraction of a
    # millimetre (21cm comes back as 21.0009cm) — hence the rounding.
    assert round(section.page_width.cm, 1) == 21
    assert round(section.page_height.cm, 1) == 29.7
    assert round(section.left_margin.cm, 1) == round(section.right_margin.cm, 1) == 1.8


def test_generate_docx_falls_back_to_the_account_name():
    document = _document(
        tailor_cv({"summary": "x", "skills": []}, []), name="Ada Lovelace", email="ada@example.com"
    )
    texts = _texts(document)

    assert "Ada Lovelace" in texts
    assert "ada@example.com" in texts


def test_generate_docx_experience_header_carries_the_date_range():
    document = _document(tailor_cv(CV, ["python"]))
    header = next(text for text in _texts(document) if text.startswith("Python Developer"))

    assert header.startswith("Python Developer — Acme")
    assert "\t2020-01 \u2013 Present" in header


def test_generate_docx_groups_skills_by_category():
    document = _document(tailor_cv(CV, ["python"]))
    texts = _texts(document)

    languages = texts.index("Languages: python")
    frameworks = texts.index("Frameworks: fastapi")
    assert languages < frameworks


def test_generate_docx_prefers_llm_skill_groups():
    """A variant that authored its own categories renders those, not the taxonomy."""
    tailored = tailor_cv(CV, ["python"])
    tailored.skill_groups = [("Backend", ["python", "fastapi"])]
    texts = _texts(_document(tailored))

    assert "Backend: python, fastapi" in texts
    assert "Languages: python" not in texts


def test_generate_docx_never_labels_a_skill_group_other():
    tailored = tailor_cv({**CV, "skills": ["python", "stakeholder management"]}, ["python"])
    assert "Additional: stakeholder management" in _texts(_document(tailored))


def test_generate_docx_labels_the_skill_group_in_bold():
    document = _document(tailor_cv(CV, ["python"]))
    group = next(p for p in document.paragraphs if p.text == "Languages: python")

    assert group.runs[0].bold is True
    assert group.runs[0].text == "Languages: "


def test_generate_docx_renders_projects_and_education():
    texts = _texts(_document(tailor_cv(CV, ["python"])))

    assert "PudimJobs" in texts
    assert "Job tracker" in texts
    assert "https://pd.dev" in texts
    assert "BSc, MIT (2015)" in texts


def test_generate_docx_omits_empty_sections():
    document = _document(tailor_cv({"summary": "x", "skills": []}, []))

    assert "Skills" not in _texts(document)
    assert "Experience" not in _texts(document)
    assert "Projects" not in _texts(document)
    assert "Summary" in _texts(document)


def test_generate_docx_stays_plain_for_ats_parsers():
    """Single column, real list styles: no tables, text boxes or images."""
    docx = generate_docx(tailor_cv(CV, ["python"]))
    document = Document(BytesIO(docx))

    assert document.tables == []
    assert "word/media" not in zipfile.ZipFile(BytesIO(docx)).namelist()
    assert b"w:txbxContent" not in zipfile.ZipFile(BytesIO(docx)).read("word/document.xml")


# --- Localised headings (ADR 024) -------------------------------------------


def test_generate_docx_uses_the_brazilian_portuguese_headings():
    document = _document(tailor_cv(CV, ["python"]), labels=labels_for("pt-BR"))
    texts = _texts(document)

    for heading in (
        "Resumo",
        "Experiência",
        "Competências",
        "Projetos",
        "Formação",
        "Certificações",
        "Idiomas",
    ):
        assert heading in texts, heading
    assert "Summary" not in texts
    # Word styles are untouched by the translation: it is still an ATS-shaped doc.
    assert _style_for(document, "Resumo") == "Heading 1"


def test_generate_docx_localises_the_open_ended_date():
    tailored = TailoredCV.from_structure(
        {
            "summary": "x",
            "experience": [
                {"company": "Acme", "title": "Dev", "start_date": "2020-01", "bullets": []}
            ],
        }
    )
    texts = _texts(_document(tailored, labels=labels_for("pt-BR")))

    assert any("2020-01" in text and "Atual" in text for text in texts)
    assert not any("Present" in text for text in texts)

