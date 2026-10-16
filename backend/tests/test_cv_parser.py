"""Unit tests for the PDF/DOCX CV parser (extraction + structuring)."""

import pytest

from app.schemas.master_cv import CVStructure
from app.services.cv_parser import (
    InvalidFileError,
    UnsupportedFileTypeError,
    extract_text,
    extract_text_from_docx,
    extract_text_from_odt,
    extract_text_from_pdf,
    parse_cv_file,
    structure_cv,
    structure_cv_with_llm,
)
from app.services.llm_config import LlmRuntimeConfig
from tests.helpers import make_docx, make_odt, make_pdf

SAMPLE_CV = """\
Jane Doe
jane@example.com

SUMMARY
Backend engineer with 6 years of experience building pipelines.

WORK EXPERIENCE
Senior Software Engineer | Acme Corp
Jan 2020 - Present
• Built async data pipelines with Python
• Led a 4-person team

Software Engineer - Globex
2017 - 2019
- Wrote REST APIs in FastAPI

EDUCATION
BSc Computer Science, MIT (2014 - 2018)

SKILLS
Python, FastAPI, PostgreSQL, Docker, AWS

PROJECTS
PudimJobs - Job application tracker
https://github.com/example/pudimjobs
• Tailors CVs to job descriptions
"""


# --- text extraction -------------------------------------------------------


def test_extract_text_from_pdf_returns_lines():
    pdf = make_pdf("SUMMARY", "Backend engineer.", "SKILLS", "Python")
    text = extract_text_from_pdf(pdf)
    assert "SUMMARY" in text
    assert "Backend engineer." in text
    assert "Python" in text


def test_extract_text_from_docx_includes_paragraphs():
    docx = make_docx("SUMMARY", "Backend engineer.")
    assert extract_text_from_docx(docx) == "SUMMARY\nBackend engineer."


def test_extract_text_from_docx_marks_list_items_as_bullets():
    # Word keeps the bullet glyph in the list numbering, not the run text.
    docx = make_docx("SKILLS", "- Python", "- FastAPI")
    assert extract_text_from_docx(docx) == "SKILLS\n\u2022 Python\n\u2022 FastAPI"


def test_extract_text_from_docx_detects_paragraph_level_numbering():
    # Exporters (e.g. LibreOffice/Markdown tooling) put ``w:numPr`` directly on
    # the paragraph instead of using a bullet style; that path must work too.
    from io import BytesIO

    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    document.add_paragraph("SKILLS")
    paragraph = document.add_paragraph("Python")
    properties = paragraph._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number_id = OxmlElement("w:numId")
    number_id.set(qn("w:val"), "1")
    numbering.append(level)
    numbering.append(number_id)
    properties.append(numbering)
    buffer = BytesIO()
    document.save(buffer)

    assert extract_text_from_docx(buffer.getvalue()) == "SKILLS\n\u2022 Python"


def test_extract_text_dispatches_by_extension():
    assert extract_text("cv.pdf", make_pdf("Hello")) == "Hello"
    assert extract_text("cv.docx", make_docx("Hello")) == "Hello"
    assert extract_text("cv.odt", make_odt("Hello")) == "Hello"


def test_extract_text_from_odt_includes_paragraphs():
    odt = make_odt("SUMMARY", "Backend engineer.", "SKILLS", "Python & FastAPI")
    assert extract_text_from_odt(odt) == ("SUMMARY\nBackend engineer.\nSKILLS\nPython & FastAPI")


def test_extract_text_from_odt_marks_list_items_as_bullets():
    # ODF keeps the bullet glyph in the list style, not the text; extraction must
    # synthesize a marker so list items are recognized as bullets downstream.
    odt = make_odt("WORK EXPERIENCE", "Engineer | Acme", "- Built pipelines", "- Led a team")
    assert extract_text_from_odt(odt) == (
        "WORK EXPERIENCE\nEngineer | Acme\n\u2022 Built pipelines\n\u2022 Led a team"
    )


def test_extract_text_from_odt_rejects_invalid_bytes():
    with pytest.raises(InvalidFileError):
        extract_text_from_odt(b"not an odt at all")


def test_extract_text_from_odt_rejects_zip_without_content():
    import zipfile
    from io import BytesIO

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
    with pytest.raises(InvalidFileError):
        extract_text_from_odt(buffer.getvalue())


def test_extract_text_rejects_unknown_extension():
    with pytest.raises(UnsupportedFileTypeError):
        extract_text("cv.txt", b"plain")


def test_extract_text_from_pdf_rejects_invalid_bytes():
    with pytest.raises(InvalidFileError):
        extract_text_from_pdf(b"not a pdf at all")


def test_extract_text_from_docx_rejects_invalid_bytes():
    with pytest.raises(InvalidFileError):
        extract_text_from_docx(b"not a docx at all")


# --- rule-based structuring ------------------------------------------------


def test_structure_cv_parses_all_sections():
    cv = structure_cv(SAMPLE_CV)
    assert "Backend engineer" in cv.summary
    assert len(cv.experience) == 2

    first = cv.experience[0]
    assert first.title == "Senior Software Engineer"
    assert first.company == "Acme Corp"
    assert first.start_date == "2020-01"
    assert first.end_date == "Present"
    assert first.bullets == [
        "Built async data pipelines with Python",
        "Led a 4-person team",
    ]

    second = cv.experience[1]
    assert second.title == "Software Engineer"
    assert second.company == "Globex"
    assert second.start_date == "2017"
    assert second.end_date == "2019"
    assert second.bullets == ["Wrote REST APIs in FastAPI"]

    assert len(cv.education) == 1
    assert cv.education[0].institution == "MIT"
    assert cv.education[0].degree == "BSc Computer Science"
    assert cv.education[0].year == "2014"

    assert cv.skills == ["Python", "FastAPI", "PostgreSQL", "Docker", "AWS"]

    assert len(cv.projects) == 1
    assert cv.projects[0].name == "PudimJobs"
    assert cv.projects[0].link == "https://github.com/example/pudimjobs"


def test_structure_cv_empty_text_is_schema_valid():
    cv = structure_cv("")
    assert cv.summary == ""
    assert cv.experience == []
    assert cv.education == []
    assert cv.skills == []
    assert cv.projects == []
    CVStructure.model_validate(cv.model_dump())


def test_structure_cv_ignores_unknown_sections():
    text = (
        "WORK EXPERIENCE\nEngineer | Acme\n2019 - 2021\n- did things\n\n"
        "LANGUAGES\nEnglish\n\n"
        "SKILLS\nPython, SQL\n"
    )
    cv = structure_cv(text)
    assert len(cv.experience) == 1
    assert cv.skills == ["Python", "SQL"]


# A skills section written the way hiring-manager-friendly CVs do it: one labelled
# line per group, multi-word and slashed skills included.
LABELLED_SKILLS_CV = """\
Jane Doe

SKILLS
Languages: C++, Rust, Python, Golang, C#, Ruby, JavaScript, TypeScript, SQL
Frameworks: Flask, Qt, Boost, Protobuf, gRPC, Apache Thrift, Tauri, RoR, Celery, n8n
Cloud & DevOps: AWS, Azure, Docker, Jenkins, Terraform, CI/CD, Artifactory, Nginx, CDK
Databases: PostgreSQL, MongoDB, Redis, SQLite, BigQuery
Concepts: Agile, Scrum, Kanban, TDD, Microservices
Tools: Linux, Windows, macOS, Git, Confluence, JIRA
Methodologies: Test-Driven Development
"""


def test_parse_skills_strips_category_labels():
    # Regression: the label words ("Frameworks", "Cloud", "DevOps", ...) used to be
    # split out of the line *as skills*, so they ended up rendered under "Additional".
    cv = structure_cv(LABELLED_SKILLS_CV)

    assert "Flask" in cv.skills
    assert "PostgreSQL" in cv.skills
    assert "Nginx" in cv.skills
    for label in (
        "Frameworks",
        "Databases",
        "Concepts",
        "Tools",
        "Methodologies",
        "Languages",
        "Cloud",
        "DevOps",
    ):
        assert label not in cv.skills


def test_parse_skills_keeps_multiword_and_slashed_skills_intact():
    # Regression: single spaces and slashes used to be delimiters, which shredded
    # these into fragments ("Apache" + "Thrift", "CI" + "CD") that no taxonomy lookup
    # could place.
    cv = structure_cv(LABELLED_SKILLS_CV)

    assert "Apache Thrift" in cv.skills
    assert "Test-Driven Development" in cv.skills
    assert "CI/CD" in cv.skills
    for fragment in ("Apache", "Thrift", "CI", "CD", "Test-Driven", "Development"):
        assert fragment not in cv.skills


def test_parse_skills_splits_on_bullets_columns_and_numbered_lines():
    cv = structure_cv(
        "SKILLS\n"
        "\u2022 Python\n"
        "\u2022 SQL\n"
        "Docker  Kubernetes\n"
        "1. Terraform\n"
        "2. AWS\n"
    )

    assert cv.skills == ["Python", "SQL", "Docker", "Kubernetes", "Terraform", "AWS"]


def test_structure_cv_keeps_a_label_only_skills_line_in_the_skills_section():
    # Regression: the bare "Languages:" line matched the spoken-languages heading, so
    # every skill after it was filed under `cv.languages` and vanished from the CV.
    cv = structure_cv("SKILLS\nLanguages:\nC++, Rust, Python\n\nTools: Git, Docker\n")

    assert cv.skills == ["C++", "Rust", "Python", "Git", "Docker"]
    assert cv.languages == []


def test_structure_cv_still_honours_a_languages_section_heading():
    cv = structure_cv("SKILLS\nPython, SQL\n\nLANGUAGES\nEnglish (fluent)\n")

    assert cv.skills == ["Python", "SQL"]
    assert cv.languages == ["English (fluent)"]


def test_structure_cv_still_switches_section_on_a_colon_terminated_heading():
    # Only *skill-group* labels are treated as labels; "PROJECTS:" still switches.
    cv = structure_cv("SKILLS\nPython\n\nPROJECTS:\nPudimJobs - Job tracker\n")

    assert cv.skills == ["Python"]
    assert [p.name for p in cv.projects] == ["PudimJobs"]


def test_structure_cv_keeps_wrapped_bullets_with_their_entry():
    # PDF text extraction yields Symbol-font bullets (\uf0b7) and wraps long
    # bullets across lines; neither should be mistaken for a new job header.
    text = (
        "WORK EXPERIENCE\n"
        "Engineer | Acme\n"
        "Jan 2020 - Present\n"
        "\uf0b7 Built async data pipelines with Python,\n"
        "using asyncio and PostgreSQL\n"
        "\uf0b7 Led a 4-person team\n"
        "Senior Engineer - Globex\n"
        "2017 - 2019\n"
        "\uf0b7 Reviewed code-\n"
        "quality reports\n"
    )
    cv = structure_cv(text)
    assert len(cv.experience) == 2

    first = cv.experience[0]
    assert first.title == "Engineer"
    assert first.company == "Acme"
    assert first.bullets == [
        "Built async data pipelines with Python, using asyncio and PostgreSQL",
        "Led a 4-person team",
    ]

    second = cv.experience[1]
    assert second.title == "Senior Engineer"
    assert second.company == "Globex"
    # Hyphenated wrap ("code-" + "quality") rejoins without a space.
    assert second.bullets == ["Reviewed code-quality reports"]


def test_structure_cv_backfills_wrapped_education_year():
    # PDF extraction wraps the year onto its own line; it should attach to the
    # preceding education entry rather than being dropped.
    text = (
        "EDUCATION\n"
        "Specialization \u2013 Computer Engineering, Universidade Municipal de "
        "Sao Caetano do Sul (USCS),\n"
        "2017\n"
    )
    cv = structure_cv(text)
    assert len(cv.education) == 1
    assert cv.education[0].year == "2017"
    assert cv.education[0].institution == "Universidade Municipal de Sao Caetano do Sul (USCS)"


def test_structure_cv_merges_multiline_entry_headers():
    # Some layouts (e.g. Word exports) split each job over two lines: title+dates
    # first, then company+location. Those must stay a single entry.
    text = (
        "WORK EXPERIENCE\n"
        "Senior Software Engineer\tJan 2020 - Present\n"
        "Acme Corp\tRemote\n"
        "\u2022 Built pipelines\n"
        "Software Engineer\t2017 - 2019\n"
        "Globex\tRemote\n"
        "\u2022 Wrote REST APIs\n"
    )
    cv = structure_cv(text)
    assert len(cv.experience) == 2
    assert cv.experience[0].start_date == "2020-01"
    assert cv.experience[0].end_date == "Present"
    assert cv.experience[0].bullets == ["Built pipelines"]
    assert cv.experience[1].start_date == "2017"
    assert cv.experience[1].end_date == "2019"
    assert cv.experience[1].bullets == ["Wrote REST APIs"]


def test_structure_cv_recovers_date_range_wrapped_across_lines():
    # Regression: PDF extraction can break a header mid-date-range
    # (the month lands on one line and the year on the next); both ends must
    # still be recovered.
    text = (
        "WORK EXPERIENCE\n"
        "Microsoft | Software Development Engineer (Sao Paulo - SP) Jan 2022 \u2013 Jul \n"
        "2022\n"
        "\u2022 Shipped dashboards\n"
    )
    cv = structure_cv(text)
    assert len(cv.experience) == 1
    assert cv.experience[0].start_date == "2022-01"
    assert cv.experience[0].end_date == "2022-07"
    assert cv.experience[0].bullets == ["Shipped dashboards"]


# --- LLM structuring (graceful degradation) --------------------------------


class _FakeResponse:
    def __init__(self, content: str):
        self._content = content
        # Read by `llm_client._rejects_json_mode` (retry-on-400 path).
        self.status_code = 200
        self.text = ""

    def raise_for_status(self) -> None:
        pass

    def json(self) -> dict:
        return {"choices": [{"message": {"content": self._content}}]}


class _FakeClient:
    """In-memory stand-in for httpx.AsyncClient (mirrors test_cv_tailor)."""

    reply = (
        '{"summary": "LLM summary", "experience": [], "education": [], '
        '"skills": ["llm-skill"], "projects": []}'
    )

    def __init__(self, *args, **kwargs):
        self.requests: list[dict] = []

    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False

    async def post(self, url, *, headers=None, json=None, **kwargs):
        self.requests.append(json)
        return _FakeResponse(self.reply)


def _llm_config() -> LlmRuntimeConfig:
    return LlmRuntimeConfig(
        enabled=True,
        api_key="sk-test",
        base_url="https://api.openai.com/v1",
        model="gpt-4o-mini",
    )


async def test_structure_cv_with_llm_parses_json(monkeypatch):
    from app.services import llm_client

    fake = _FakeClient()
    monkeypatch.setattr(llm_client.httpx, "AsyncClient", lambda *a, **k: fake)

    result = await structure_cv_with_llm("raw text", _llm_config())
    assert result is not None
    assert result.summary == "LLM summary"
    assert result.skills == ["llm-skill"]
    assert fake.requests[0]["messages"][0]["role"] == "system"


async def test_structure_cv_with_llm_returns_none_on_garbage(monkeypatch):
    from app.services import llm_client

    fake = _FakeClient()
    fake.reply = "not json at all"
    monkeypatch.setattr(llm_client.httpx, "AsyncClient", lambda *a, **k: fake)

    assert await structure_cv_with_llm("raw text", _llm_config()) is None


async def test_structure_cv_with_llm_returns_none_on_invalid_shape(monkeypatch):
    from app.services import llm_client

    fake = _FakeClient()
    fake.reply = '{"summary": 123}'
    monkeypatch.setattr(llm_client.httpx, "AsyncClient", lambda *a, **k: fake)

    assert await structure_cv_with_llm("raw text", _llm_config()) is None


# --- end-to-end parse_cv_file ----------------------------------------------


async def test_parse_cv_file_uses_llm_when_configured(monkeypatch):
    from app.services import llm_client

    fake = _FakeClient()
    monkeypatch.setattr(llm_client.httpx, "AsyncClient", lambda *a, **k: fake)

    docx = make_docx("WORK EXPERIENCE", "Engineer | Acme", "Jan 2020 - Present", "- did things")
    result = await parse_cv_file("cv.docx", docx, llm_config=_llm_config())
    assert result.summary == "LLM summary"
    assert result.skills == ["llm-skill"]


async def test_parse_cv_file_falls_back_to_rule_based_when_llm_fails(monkeypatch):
    from app.services import llm_client

    fake = _FakeClient()
    fake.reply = "boom"
    monkeypatch.setattr(llm_client.httpx, "AsyncClient", lambda *a, **k: fake)

    docx = make_docx("WORK EXPERIENCE", "Engineer | Acme", "Jan 2020 - Present", "- did things")
    result = await parse_cv_file("cv.docx", docx, llm_config=_llm_config())
    assert result.experience[0].title == "Engineer"
    assert result.experience[0].company == "Acme"
    assert result.experience[0].start_date == "2020-01"
    assert result.experience[0].end_date == "Present"


async def test_parse_cv_file_rule_based_without_llm_config():
    docx = make_docx("WORK EXPERIENCE", "Engineer | Acme", "Jan 2020 - Present", "- did things")
    result = await parse_cv_file("cv.docx", docx)
    assert result.experience[0].title == "Engineer"
    assert result.experience[0].company == "Acme"


async def test_parse_cv_file_pdf_end_to_end():
    pdf = make_pdf("WORK EXPERIENCE", "Engineer | Acme", "2020 - Present", "- did things")
    result = await parse_cv_file("cv.pdf", pdf)
    assert result.experience[0].title == "Engineer"
    assert result.experience[0].company == "Acme"
    assert result.experience[0].start_date == "2020"
    assert result.experience[0].end_date == "Present"


async def test_parse_cv_file_odt_end_to_end():
    odt = make_odt(
        "SUMMARY",
        "Backend engineer.",
        "WORK EXPERIENCE",
        "Engineer | Acme",
        "Jan 2020 - Present",
        "- did things",
        "SKILLS",
        "Python, FastAPI",
    )
    result = await parse_cv_file("cv.odt", odt)
    assert "Backend engineer." in result.summary
    assert result.experience[0].title == "Engineer"
    assert result.experience[0].company == "Acme"
    assert result.experience[0].start_date == "2020-01"
    assert result.experience[0].end_date == "Present"
    assert "Python" in result.skills


async def test_parse_cv_file_odt_list_bullets_stay_with_entry():
    # Regression: list-item bullets used to be dropped at extraction, so every
    # bullet line became its own (bogus) experience entry.
    odt = make_odt(
        "WORK EXPERIENCE",
        "Engineer | Acme",
        "Jan 2020 - Present",
        "- Built pipelines",
        "- Led a team",
    )
    result = await parse_cv_file("cv.odt", odt)
    assert len(result.experience) == 1
    assert result.experience[0].bullets == ["Built pipelines", "Led a team"]


# --- richer structure: contact, certifications, languages -------------------


def test_structure_cv_extracts_contact_block():
    cv = structure_cv(SAMPLE_CV)
    assert cv.full_name == "Jane Doe"
    assert cv.email == "jane@example.com"


def test_structure_cv_extracts_phone_location_and_links():
    text = (
        "Jane Doe\n"
        "jane@example.com | +55 11 91234-5678 | Sao Paulo, Brazil\n"
        "https://github.com/janedoe\n"
        "SUMMARY\n"
        "Backend engineer\n"
    )
    cv = structure_cv(text)
    assert cv.full_name == "Jane Doe"
    assert cv.email == "jane@example.com"
    assert cv.phone == "+55 11 91234-5678"
    assert cv.location == "Sao Paulo, Brazil"
    assert cv.links == ["https://github.com/janedoe"]


def test_structure_cv_parses_certifications_and_languages():
    text = (
        "CERTIFICATIONS\n"
        "AWS Certified Solutions Architect, CKA: Certified Kubernetes Administrator\n"
        "LANGUAGES\n"
        "English (fluent), Portuguese (native)\n"
    )
    cv = structure_cv(text)
    assert cv.certifications == [
        "AWS Certified Solutions Architect",
        "CKA: Certified Kubernetes Administrator",
    ]
    assert cv.languages == ["English (fluent)", "Portuguese (native)"]


def test_structure_cv_defaults_contact_when_absent():
    cv = structure_cv("SUMMARY\nBackend engineer\n")
    assert cv.full_name == ""
    assert cv.email == ""
    assert cv.phone == ""
    assert cv.location == ""
    assert cv.links == []
    assert cv.certifications == []
    assert cv.languages == []
