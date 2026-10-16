"""Tests for the master CV API (versioning + audit)."""

import json
import uuid
from io import BytesIO

from docx import Document
from sqlalchemy import select
from structlog.testing import capture_logs

from app.models.audit_log import AuditLog
from app.models.generated_cv import GeneratedCV
from app.routers import cv as cv_router
from app.routers.cv import DOCX_MEDIA_TYPE
from app.services import cv_translator
from app.services.llm_config import LlmRuntimeConfig
from tests.helpers import create_user

CV_PAYLOAD = {
    "summary": "Backend engineer",
    "experience": [
        {
            "company": "Acme",
            "title": "Engineer",
            "start_date": "2020-01",
            "bullets": ["Built async pipelines"],
        }
    ],
    "education": [{"institution": "MIT", "degree": "BSc Computer Science"}],
    "skills": ["python", "fastapi"],
    "projects": [{"name": "PudimJobs", "description": "Job application tracker"}],
}


async def test_create_cv_versions_and_current(auth_client):
    client, _, _ = auth_client

    first = await client.post("/api/cv", json={"structured_json": CV_PAYLOAD})
    assert first.status_code == 201
    assert first.json()["version"] == 1
    assert first.json()["is_current"] is True
    assert first.json()["label"] == "CV v1"

    second = await client.post("/api/cv", json={"structured_json": CV_PAYLOAD, "label": "Tailored"})
    assert second.status_code == 201
    assert second.json()["version"] == 2
    assert second.json()["label"] == "Tailored"
    assert second.json()["is_current"] is True

    current = await client.get("/api/cv/current")
    assert current.json()["version"] == 2

    versions = await client.get("/api/cv")
    assert [v["version"] for v in versions.json()] == [2, 1]
    assert versions.json()[1]["is_current"] is False


async def test_cv_creation_writes_audit_log(auth_client, db_session):
    client, user, _ = auth_client
    await client.post("/api/cv", json={"structured_json": CV_PAYLOAD})

    result = await db_session.execute(select(AuditLog).where(AuditLog.user_id == user.id))
    entries = result.scalars().all()
    assert len(entries) == 1
    assert entries[0].entity_type == "master_cv"
    assert entries[0].action == "created"
    assert entries[0].changes == {"version": 1}


async def test_export_cv_pdf(auth_client):
    client, _, _ = auth_client
    response = await client.post("/api/cv/pdf", json=CV_PAYLOAD)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


async def test_export_cv_docx(auth_client):
    """`format=docx` returns a real Word document from the same endpoint."""
    client, _, _ = auth_client
    response = await client.post("/api/cv/pdf", json=CV_PAYLOAD, params={"format": "docx"})

    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_MEDIA_TYPE
    assert response.headers["content-disposition"] == 'attachment; filename="cv.docx"'
    assert response.content.startswith(b"PK")
    Document(BytesIO(response.content))  # parses as a Word document


async def test_export_cv_docx_keeps_llm_skill_groups(auth_client):
    """The editor export used to drop `skill_groups`; the DOCX must show them."""
    client, _, _ = auth_client
    payload = {
        **CV_PAYLOAD,
        "skill_groups": [{"label": "Backend", "items": ["python", "fastapi"]}],
    }
    response = await client.post("/api/cv/pdf", json=payload, params={"format": "docx"})

    texts = [paragraph.text for paragraph in Document(BytesIO(response.content)).paragraphs]
    assert "Backend: python, fastapi" in texts
    assert not any(text.startswith("Languages:") for text in texts)


async def test_export_cv_rejects_unknown_format(auth_client):
    client, _, _ = auth_client
    response = await client.post("/api/cv/pdf", json=CV_PAYLOAD, params={"format": "odt"})

    assert response.status_code == 400
    assert "odt" in response.json()["detail"]


async def test_export_cv_rejects_unknown_template(auth_client):
    client, _, _ = auth_client
    response = await client.post("/api/cv/pdf", json=CV_PAYLOAD, params={"template": "fancy"})
    assert response.status_code == 400


async def test_cv_requires_auth(client):
    assert (await client.get("/api/cv")).status_code == 401


async def test_delete_current_cv_promotes_next_version(auth_client):
    client, _, _ = auth_client
    first = (await client.post("/api/cv", json={"structured_json": CV_PAYLOAD})).json()
    second = (
        await client.post("/api/cv", json={"structured_json": CV_PAYLOAD, "label": "Tailored"})
    ).json()

    response = await client.delete(f"/api/cv/{second['id']}")
    assert response.status_code == 204

    current = await client.get("/api/cv/current")
    assert current.status_code == 200
    assert current.json()["id"] == first["id"]
    assert current.json()["is_current"] is True

    versions = await client.get("/api/cv")
    assert len(versions.json()) == 1


async def test_delete_non_current_cv_keeps_current(auth_client):
    client, _, _ = auth_client
    first = (await client.post("/api/cv", json={"structured_json": CV_PAYLOAD})).json()
    second = (
        await client.post("/api/cv", json={"structured_json": CV_PAYLOAD, "label": "Tailored"})
    ).json()

    response = await client.delete(f"/api/cv/{first['id']}")
    assert response.status_code == 204

    current = await client.get("/api/cv/current")
    assert current.status_code == 200
    assert current.json()["id"] == second["id"]


async def test_delete_last_cv_leaves_no_current(auth_client):
    client, _, _ = auth_client
    created = (await client.post("/api/cv", json={"structured_json": CV_PAYLOAD})).json()

    response = await client.delete(f"/api/cv/{created['id']}")
    assert response.status_code == 204

    assert (await client.get("/api/cv/current")).status_code == 404
    assert (await client.get("/api/cv")).json() == []


async def test_delete_cv_version_unknown_404(auth_client):
    client, _, _ = auth_client
    response = await client.delete("/api/cv/00000000-0000-0000-0000-000000000000")
    assert response.status_code == 404


async def test_delete_generated_cv(auth_client, db_session):
    client, user, _ = auth_client
    job = (await client.post("/api/jobs", json={"title": "Engineer", "company": "Acme"})).json()
    generated = GeneratedCV(
        user_id=user.id,
        job_id=uuid.UUID(job["id"]),
        pdf=b"%PDF-1.4 fake",
        structured_json=CV_PAYLOAD,
    )
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)

    listed = await client.get("/api/cv/generated")
    assert len(listed.json()) == 1

    response = await client.delete(f"/api/cv/generated/{generated.id}")
    assert response.status_code == 204

    assert (await client.get("/api/cv/generated")).json() == []


async def test_delete_generated_cv_unknown_404(auth_client):
    client, _, _ = auth_client
    response = await client.delete("/api/cv/generated/00000000-0000-0000-0000-000000000000")
    assert response.status_code == 404


async def test_list_generated_filters_by_job(auth_client, db_session):
    """``?job_id=`` scopes the list — this powers the application CV panel."""
    client, user, _ = auth_client
    job_a = (await client.post("/api/jobs", json={"title": "A", "company": "Acme"})).json()
    job_b = (await client.post("/api/jobs", json={"title": "B", "company": "Acme"})).json()

    db_session.add_all(
        [
            GeneratedCV(
                user_id=user.id,
                job_id=uuid.UUID(job_a["id"]),
                pdf=b"%PDF-1.4 a",
                structured_json=CV_PAYLOAD,
            ),
            GeneratedCV(
                user_id=user.id,
                job_id=uuid.UUID(job_b["id"]),
                pdf=b"%PDF-1.4 b",
                structured_json=CV_PAYLOAD,
            ),
        ]
    )
    await db_session.commit()

    unfiltered = await client.get("/api/cv/generated")
    assert len(unfiltered.json()) == 2
    assert {row["variant"] for row in unfiltered.json()} == {"rule"}

    scoped = await client.get("/api/cv/generated", params={"job_id": job_a["id"]})
    assert [row["job_id"] for row in scoped.json()] == [job_a["id"]]


async def test_list_generated_hides_content_less_artifacts(auth_client, db_session):
    """A row with no content is not a CV, so the panel must not list it (ADR 016).

    Pre-ADR-016 rows have ``structured_json`` NULL: nothing to preview, edit or
    score. Listing one let the application CV panel claim the job had a CV while
    blaming the missing ``llm`` variant for the empty body.
    """
    client, user, _ = auth_client
    job = (await client.post("/api/jobs", json={"title": "A", "company": "Acme"})).json()
    empty = GeneratedCV(
        user_id=user.id,
        job_id=uuid.UUID(job["id"]),
        pdf=b"%PDF-1.4 legacy",
    )
    usable = GeneratedCV(
        user_id=user.id,
        job_id=uuid.UUID(job["id"]),
        pdf=b"%PDF-1.4 tailored",
        structured_json=CV_PAYLOAD,
    )
    db_session.add_all([empty, usable])
    await db_session.commit()
    await db_session.refresh(empty)
    await db_session.refresh(usable)

    with capture_logs() as logs:
        response = await client.get("/api/cv/generated", params={"job_id": job["id"]})

    assert response.status_code == 200
    assert [row["id"] for row in response.json()] == [str(usable.id)]
    assert any(
        entry["event"] == "cv_artifact_without_content_hidden"
        and entry["generated_cv_ids"] == [str(empty.id)]
        and entry["variants"] == ["rule"]
        for entry in logs
    )

    # The row is kept (the next tailor run repairs it in place, keeping the
    # application link), so a direct fetch — the editor, the PDF — still works.
    detail = await client.get(f"/api/cv/generated/{empty.id}")
    assert detail.status_code == 200
    assert detail.json()["structured_json"] is None


async def test_update_generated_cv_edits_in_place(auth_client, db_session):
    """Editing a tailored CV re-renders its PDF and clears the stale ATS cache."""
    client, user, _ = auth_client
    job = (await client.post("/api/jobs", json={"title": "Engineer", "company": "Acme"})).json()
    generated = GeneratedCV(
        user_id=user.id,
        job_id=uuid.UUID(job["id"]),
        pdf=b"%PDF-1.4 old",
        structured_json={"summary": "old"},
        ats_score=42,
        ats_report={"score": 42},
    )
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)

    response = await client.put(
        f"/api/cv/generated/{generated.id}",
        json={"structured_json": {**CV_PAYLOAD, "summary": "Rewritten summary"}},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["structured_json"]["summary"] == "Rewritten summary"
    assert body["variant"] == "rule"
    assert body["ats_score"] is None
    assert body["ats_report"] is None
    assert body["job_title"] == "Engineer"

    await db_session.refresh(generated)
    assert generated.pdf[:4] == b"%PDF"

    entries = (
        (await db_session.execute(select(AuditLog).where(AuditLog.entity_type == "generated_cv")))
        .scalars()
        .all()
    )
    assert [entry.action for entry in entries] == ["updated"]


async def test_update_generated_cv_rejects_another_users_cv(auth_client, db_session):
    client, _, _ = auth_client
    other = await create_user(db_session, email="generated-owner@example.com")
    generated = GeneratedCV(user_id=other.id, pdf=b"%PDF-1.4 stub")
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)

    response = await client.put(
        f"/api/cv/generated/{generated.id}", json={"structured_json": CV_PAYLOAD}
    )

    assert response.status_code == 404


async def test_update_generated_cv_unknown_404(auth_client):
    client, _, _ = auth_client
    response = await client.put(
        "/api/cv/generated/00000000-0000-0000-0000-000000000000",
        json={"structured_json": CV_PAYLOAD},
    )
    assert response.status_code == 404


async def test_download_generated_cv_as_docx(auth_client, db_session):
    """DOCX is rendered on demand from the artifact's own structure (ADR 017)."""
    client, user, _ = auth_client
    generated = GeneratedCV(
        user_id=user.id,
        pdf=b"%PDF-1.4 stub",
        structured_json={**CV_PAYLOAD, "full_name": "Jane Doe"},
    )
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)

    response = await client.get(
        f"/api/cv/generated/{generated.id}/pdf", params={"format": "docx"}
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_MEDIA_TYPE
    assert response.headers["content-disposition"] == (
        f'attachment; filename="tailored-cv-{generated.id}.docx"'
    )
    texts = [paragraph.text for paragraph in Document(BytesIO(response.content)).paragraphs]
    assert "Jane Doe" in texts
    assert "Experience" in texts


async def test_download_generated_cv_serves_the_stored_pdf(auth_client, db_session):
    """The default (and `format=pdf`) stays the stored bytes, untouched."""
    client, user, _ = auth_client
    generated = GeneratedCV(
        user_id=user.id, pdf=b"%PDF-1.4 stored", structured_json=CV_PAYLOAD
    )
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)

    response = await client.get(f"/api/cv/generated/{generated.id}/pdf")

    assert response.content == b"%PDF-1.4 stored"
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="tailored-cv-{generated.id}.pdf"'
    )

    explicit = await client.get(
        f"/api/cv/generated/{generated.id}/pdf", params={"format": "pdf"}
    )
    assert explicit.content == b"%PDF-1.4 stored"


async def test_download_generated_cv_rejects_unknown_format(auth_client, db_session):
    client, user, _ = auth_client
    generated = GeneratedCV(user_id=user.id, pdf=b"%PDF-1.4 stub", structured_json=CV_PAYLOAD)
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)

    response = await client.get(
        f"/api/cv/generated/{generated.id}/pdf", params={"format": "odt"}
    )
    assert response.status_code == 400


async def test_download_generated_cv_docx_without_structure_404(auth_client, db_session):
    """A content-less artifact (imported, never rendered) has nothing to re-render."""
    client, user, _ = auth_client
    generated = GeneratedCV(user_id=user.id, pdf=b"%PDF-1.4 stub", structured_json=None)
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)

    response = await client.get(
        f"/api/cv/generated/{generated.id}/pdf", params={"format": "docx"}
    )
    assert response.status_code == 404


async def test_download_generated_cv_docx_rejects_another_users_artifact(auth_client, db_session):
    client, _, _ = auth_client
    other = await create_user(db_session, email="docx-owner@example.com")
    generated = GeneratedCV(user_id=other.id, pdf=b"%PDF-1.4 stub", structured_json=CV_PAYLOAD)
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)

    response = await client.get(
        f"/api/cv/generated/{generated.id}/pdf", params={"format": "docx"}
    )
    assert response.status_code == 404


# --- PT-BR export (ADR 024) -------------------------------------------------


def _enable_llm(monkeypatch, *, enabled: bool = True, api_key: str = "sk-x") -> None:
    """Point the router at a canned LLM config (the DB row is admin-owned)."""

    async def _config(_db):
        return LlmRuntimeConfig(
            enabled=enabled, api_key=api_key, base_url="http://llm", model="gpt-4o-mini"
        )

    monkeypatch.setattr(cv_router, "get_llm_config", _config)


# The taxonomy's group labels as a translated reply would carry them. The stub
# below reads the structure it was handed (a real model echoes it back) and
# translates the labels only — the chrome/data split the prompt asks for.
GROUP_LABELS_PT_BR = {
    "Languages": "Linguagens",
    "Frameworks": "Frameworks",
    "Databases": "Bancos de dados",
    "Cloud & DevOps": "Cloud & DevOps",
    "Concepts": "Conceitos",
    "Tools": "Ferramentas",
    "Additional": "Adicionais",
}


def _canned_translation(monkeypatch, **overrides) -> list[dict]:
    """Stub the service's ``chat_json``; returns the list of calls it received."""
    calls: list[dict] = []

    async def _fake_chat_json(config, **kwargs):
        calls.append(kwargs)
        sent = json.loads(kwargs["user"].split("\n\n", 1)[1])
        groups = [
            {
                "label": GROUP_LABELS_PT_BR.get(group["label"], group["label"]),
                "items": group["items"],
            }
            for group in sent.get("skill_groups") or []
        ]
        # `full_name` is identity, not prose: the reply echoes it untouched.
        return {
            **CV_PAYLOAD,
            "full_name": "Jane Doe",
            "summary": "Engenheiro backend",
            "skill_groups": groups,
            **overrides,
        }

    monkeypatch.setattr(cv_translator, "chat_json", _fake_chat_json)
    return calls


async def test_cv_capabilities_report_the_export_languages(auth_client):
    """The client cannot read admin settings, so this is how it gates PT-BR."""
    client, _, _ = auth_client
    response = await client.get("/api/cv/capabilities")

    assert response.status_code == 200
    assert response.json() == {
        "llm_enabled": False,
        "export_languages": ["en", "pt-BR"],
        "default_language": "en",
    }


async def test_cv_capabilities_follow_the_llm_config(auth_client, monkeypatch):
    client, _, _ = auth_client
    _enable_llm(monkeypatch)

    assert (await client.get("/api/cv/capabilities")).json()["llm_enabled"] is True


async def test_cv_capabilities_require_a_key_not_just_enabled(auth_client, monkeypatch):
    client, _, _ = auth_client
    _enable_llm(monkeypatch, api_key="")

    assert (await client.get("/api/cv/capabilities")).json()["llm_enabled"] is False


async def test_export_cv_in_pt_br_requires_the_llm(auth_client):
    client, _, _ = auth_client
    response = await client.post("/api/cv/pdf", json=CV_PAYLOAD, params={"language": "pt-BR"})

    assert response.status_code == 409
    assert "pt-BR" in response.json()["detail"]


async def test_export_cv_rejects_an_unknown_language(auth_client):
    client, _, _ = auth_client
    response = await client.post("/api/cv/pdf", json=CV_PAYLOAD, params={"language": "de"})

    assert response.status_code == 400
    assert "de" in response.json()["detail"]


async def test_export_cv_translates_on_demand_for_pt_br(auth_client, monkeypatch):
    """The rendered DOCX carries translated content, Portuguese headings *and* the
    skill-group labels the taxonomy would otherwise print in English."""
    client, _, _ = auth_client
    _enable_llm(monkeypatch)
    calls = _canned_translation(monkeypatch)

    response = await client.post(
        "/api/cv/pdf", json=CV_PAYLOAD, params={"format": "docx", "language": "pt-BR"}
    )

    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="cv-pt-br.docx"'
    texts = [paragraph.text for paragraph in Document(BytesIO(response.content)).paragraphs]
    assert "Engenheiro backend" in texts
    assert "Resumo" in texts
    assert "Experiência" in texts
    # CV_PAYLOAD stores no `skill_groups`, so the renderer groups by taxonomy —
    # those labels are what a PT-BR CV used to show in English.
    assert "Linguagens: python" in texts
    assert "Frameworks: fastapi" in texts
    assert "Summary" not in texts
    assert not any(text.startswith("Languages:") for text in texts)
    assert len(calls) == 1
    assert "Brazilian Portuguese" in calls[0]["system"]


async def test_export_cv_fails_when_the_translation_fails(auth_client, monkeypatch):
    """A broken translation is a 502, never an English CV with PT headings."""
    client, _, _ = auth_client
    _enable_llm(monkeypatch)
    _canned_translation(monkeypatch, summary="")

    response = await client.post("/api/cv/pdf", json=CV_PAYLOAD, params={"language": "pt-BR"})

    assert response.status_code == 502


async def test_export_cv_fails_when_the_translation_flattens_the_skill_groups(
    auth_client, monkeypatch
):
    """The groups handed to the model are content: losing them fails the export
    rather than rendering a CV with an English skills heading."""
    client, _, _ = auth_client
    _enable_llm(monkeypatch)
    _canned_translation(monkeypatch, skill_groups=[])

    response = await client.post("/api/cv/pdf", json=CV_PAYLOAD, params={"language": "pt-BR"})

    assert response.status_code == 502




async def test_download_generated_cv_in_pt_br_renders_on_demand(
    auth_client, db_session, monkeypatch
):
    """The stored PDF is English, so PT-BR re-renders it from the structure."""
    client, user, _ = auth_client
    generated = GeneratedCV(
        user_id=user.id,
        pdf=b"%PDF-1.4 stored",
        structured_json={**CV_PAYLOAD, "full_name": "Jane Doe"},
    )
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)
    _enable_llm(monkeypatch)
    calls = _canned_translation(monkeypatch)

    response = await client.get(
        f"/api/cv/generated/{generated.id}/pdf", params={"language": "pt-BR"}
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="tailored-cv-{generated.id}-pt-br.pdf"'
    )
    assert response.content != b"%PDF-1.4 stored"
    assert response.content.startswith(b"%PDF")
    assert len(calls) == 1


async def test_download_generated_cv_in_pt_br_requires_the_llm(auth_client, db_session):
    client, user, _ = auth_client
    generated = GeneratedCV(user_id=user.id, pdf=b"%PDF-1.4 stored", structured_json=CV_PAYLOAD)
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)

    response = await client.get(
        f"/api/cv/generated/{generated.id}/pdf", params={"language": "pt-BR"}
    )

    assert response.status_code == 409


async def test_download_generated_cv_pt_br_without_a_structure_is_a_404(auth_client, db_session):
    client, user, _ = auth_client
    generated = GeneratedCV(user_id=user.id, pdf=b"%PDF-1.4 stored", structured_json=None)
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)

    response = await client.get(
        f"/api/cv/generated/{generated.id}/pdf", params={"language": "pt-BR"}
    )

    assert response.status_code == 404


async def test_download_generated_cv_pt_br_docx_carries_portuguese_headings(
    auth_client, db_session, monkeypatch
):
    client, user, _ = auth_client
    generated = GeneratedCV(
        user_id=user.id, pdf=b"%PDF-1.4 stored", structured_json={**CV_PAYLOAD, "full_name": "J"}
    )
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)
    _enable_llm(monkeypatch)
    _canned_translation(monkeypatch)

    response = await client.get(
        f"/api/cv/generated/{generated.id}/pdf",
        params={"format": "docx", "language": "pt-BR"},
    )

    assert response.status_code == 200
    assert response.headers["content-disposition"] == (
        f'attachment; filename="tailored-cv-{generated.id}-pt-br.docx"'
    )
    texts = [paragraph.text for paragraph in Document(BytesIO(response.content)).paragraphs]
    assert "Resumo" in texts
    assert "Engenheiro backend" in texts
    assert "Summary" not in texts


async def test_download_generated_cv_rejects_an_unknown_language(auth_client, db_session):
    client, user, _ = auth_client
    generated = GeneratedCV(user_id=user.id, pdf=b"%PDF-1.4 stored", structured_json=CV_PAYLOAD)
    db_session.add(generated)
    await db_session.commit()
    await db_session.refresh(generated)

    response = await client.get(
        f"/api/cv/generated/{generated.id}/pdf", params={"language": "pt"}
    )

    assert response.status_code == 400

