"""DOCX export: renders a tailored CV with python-docx.

DOCX is the *editable* sibling of the canonical PDF. Application forms and
recruiters routinely ask for a Word file, and it is the format ATS parsers read
most reliably — as long as it stays plain.

The document is built from the same ``TailoredCV`` the PDF renderer walks, not
from the HTML templates: those are WeasyPrint/CSS-shaped, and HTML→DOCX
conversion flattens them (a two-column block collapses into one paragraph,
hand-drawn bullets lose their list semantics). The trade-off is that this export
is *content-identical* to the PDF, never pixel-identical: Word has its own style
vocabulary and no equivalent of the template's rules, colours or letter-spacing.

Everything stays in a single column using real Word styles (``Title``,
``Heading 1``, ``List Bullet``): no tables, no text boxes, no images.
"""

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Cm, Emu, Pt

from app.services.cv_tailor import TailoredCV

# Same separator the PDF contact line uses between fields.
_CONTACT_SEPARATOR = " • "
# Contact fields render at 8pt, the size at which the PDF header fits a
# five-field line on one row (ADR 021).
_CONTACT_FONT_SIZE = 8
# A4 with the PDF template's 1.8cm side margins. python-docx defaults to Letter
# with 1.25in margins — only 432pt of usable width — where even the 8pt contact
# line overflows and wraps.
_A4_PAGE = (Cm(21), Cm(29.7))
_SIDE_MARGIN = Cm(1.8)
# Typographic en dash, escaped because a literal one trips RUF001 (ambiguous
# unicode). It is the separator the PDF's experience header uses.
_DATE_SEPARATOR = "\u2013"


def _set_a4_page(document: Document) -> None:
    """Give the document the same page box the PDF template renders on (ADR 021)."""
    section = document.sections[0]
    section.page_width, section.page_height = _A4_PAGE
    section.left_margin = section.right_margin = _SIDE_MARGIN


def _usable_width(document: Document) -> int:
    """Text width in EMU, so a right-aligned tab stop lands flush with the margin."""
    section = document.sections[0]
    return int(section.page_width - section.left_margin - section.right_margin)


def _add_dated_line(document: Document, left: str, right: str = "") -> None:
    """A bold left-hand line with an optional right-aligned, smaller date."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Emu(_usable_width(document)), WD_TAB_ALIGNMENT.RIGHT
    )
    paragraph.add_run(left).bold = True
    if right:
        paragraph.add_run(f"\t{right}").font.size = Pt(9)


def _date_range(item: dict) -> str:
    """The ``start``-``end`` range shown in the PDF's experience header."""
    start = str(item.get("start_date") or "").strip()
    end = str(item.get("end_date") or "").strip()
    if not start and not end:
        return ""
    return f"{start} {_DATE_SEPARATOR} {end or 'Present'}"


def _add_bullets(document: Document, bullets: list[object]) -> None:
    for bullet in bullets:
        text = str(bullet).strip()
        if text:
            document.add_paragraph(text, style="List Bullet")


def _add_list_section(document: Document, heading: str, items: list[str]) -> None:
    values = [str(item).strip() for item in items if str(item).strip()]
    if not values:
        return
    document.add_heading(heading, level=1)
    for value in values:
        document.add_paragraph(value)


def _add_header(document: Document, tailored: TailoredCV, name: str, email: str) -> None:
    """Name + one contact line holding every field, as the PDF header draws it."""
    display_name = tailored.full_name or name
    if display_name:
        paragraph = document.add_paragraph(display_name, style="Title")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    fields = [
        value for value in (tailored.email or email, tailored.phone, tailored.location) if value
    ]
    fields.extend(tailored.links)
    if fields:
        # One line, not a details paragraph plus a links paragraph: Word wraps it
        # between fields when it has to, and the wider A4 page box means it
        # usually does not (ADR 021).
        paragraph = document.add_paragraph(_CONTACT_SEPARATOR.join(fields))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.size = Pt(_CONTACT_FONT_SIZE)


def _add_experience(document: Document, experience: list[dict]) -> None:
    if not experience:
        return
    document.add_heading("Experience", level=1)
    for item in experience:
        if not isinstance(item, dict):
            continue
        heading = " — ".join(
            part
            for part in (
                str(item.get("title") or "").strip(),
                str(item.get("company") or "").strip(),
            )
            if part
        )
        date = _date_range(item)
        if heading:
            _add_dated_line(document, heading, date)
        elif date:
            document.add_paragraph(date)
        _add_bullets(document, item.get("bullets") or [])


def _add_skills(document: Document, groups: list[tuple[str, list[str]]]) -> None:
    """One bold ``Label:`` line per category, exactly like the PDF section."""
    groups = [(label, items) for label, items in groups if items]
    if not groups:
        return
    document.add_heading("Skills", level=1)
    for label, items in groups:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(", ".join(items))


def _add_projects(document: Document, projects: list[dict]) -> None:
    if not projects:
        return
    document.add_heading("Projects", level=1)
    for item in projects:
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or "").strip()
        if name:
            _add_dated_line(document, name)
        description = str(item.get("description") or "").strip()
        if description:
            document.add_paragraph(description)
        link = str(item.get("link") or "").strip()
        if link:
            paragraph = document.add_paragraph(link)
            for run in paragraph.runs:
                run.font.size = Pt(9)


def _add_education(document: Document, education: list[dict]) -> None:
    if not education:
        return
    document.add_heading("Education", level=1)
    for item in education:
        if not isinstance(item, dict):
            continue
        line = ", ".join(
            part
            for part in (
                str(item.get("degree") or "").strip(),
                str(item.get("institution") or "").strip(),
            )
            if part
        )
        year = str(item.get("year") or "").strip()
        if year:
            line = f"{line} ({year})" if line else year
        if line:
            document.add_paragraph(line)


def generate_docx(
    tailored: TailoredCV,
    *,
    name: str = "",
    email: str = "",
) -> bytes:
    """Render a tailored CV as a Word document.

    ``name``/``email`` are the account-derived fallbacks for a CV that carries no
    contact header (same contract as ``generate_pdf``).
    """
    document = Document()
    _set_a4_page(document)
    _add_header(document, tailored, name=name, email=email)
    if tailored.summary:
        document.add_heading("Summary", level=1)
        document.add_paragraph(tailored.summary)
    _add_experience(document, tailored.experience)
    _add_skills(document, tailored.resolved_skill_groups())
    _add_projects(document, tailored.projects)
    _add_education(document, tailored.education)
    _add_list_section(document, "Certifications", tailored.certifications)
    _add_list_section(document, "Languages", tailored.languages)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()

